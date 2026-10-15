import zipfile, sys
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

d = "tests/samples/cleaned/"
Document(d + "report.docx"); Presentation(d + "slides.pptx"); load_workbook(d + "table.xlsx")
print("все три файла открываются библиотеками")

for f in ["report.docx", "slides.pptx", "table.xlsx"]:
    z = zipfile.ZipFile(d + f)
    names = [i.filename for i in z.infolist()]
    dates = {i.date_time for i in z.infolist()}
    ct = z.read("[Content_Types].xml").decode()
    rels = z.read("_rels/.rels").decode()
    print(f"\n{f}: первые части {names[:2]}, даты внутри {dates}")
    print("  custom.xml:", "docProps/custom.xml" in names, "| в типах:", "custom.xml" in ct, "| в связях:", "custom-properties" in rels)
    print("  core.xml:", z.read("docProps/core.xml").decode()[:400])

z = zipfile.ZipFile(d + "report.docx")
png = [n for n in z.namelist() if n.endswith(".png")][0]
b = z.read(png)
print("\nPNG: tEXt" , b"tEXt" in b, "caBX", b"caBX" in b, "matplotlib", b"matplotlib" in b)
rels = z.read("word/_rels/document.xml.rels").decode()
print("ссылка:", [x for x in rels.split('"') if "cbu.uz" in x])
print("автор примечания:", [x for x in z.read("word/comments.xml").decode().split(" ") if "author" in x or "initials" in x])

z = zipfile.ZipFile(d + "slides.pptx")
jpg = [n for n in z.namelist() if n.endswith((".jpg", ".jpeg"))][0]
b = z.read(jpg)
print("JPEG: Exif", b"Exif" in b, "DALL", b"DALL" in b, "comment", b"Generated" in b)
emb = [n for n in z.namelist() if "embeddings" in n][0]
import io
ez = zipfile.ZipFile(io.BytesIO(z.read(emb)))
print("встроенный xlsx core:", ez.read("docProps/core.xml").decode()[150:450])

z = zipfile.ZipFile(d + "table.xlsx")
print("xlsx comments:", [n for n in z.namelist() if "comment" in n], z.read([n for n in z.namelist() if n.startswith("xl/comments")][0]).decode()[:300])
